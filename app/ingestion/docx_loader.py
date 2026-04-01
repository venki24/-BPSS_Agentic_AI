"""
DOCX ingestion: load → chunk → tag metadata → return Documents.
"""
import re
from pathlib import Path
from typing import List

from docx import Document as DocxDoc
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings
from app.observability import logger

_CAND_RE = re.compile(r"CAND-\d{3}", re.IGNORECASE)


def _extract_candidate_id(text: str) -> str:
    m = _CAND_RE.search(text)
    return m.group(0).upper() if m else ""


def _read_docx(path: Path) -> str:
    doc = DocxDoc(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Also capture table content
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n\n".join(paragraphs)


def load_docx(file_path: str | Path) -> List[Document]:
    path = Path(file_path)
    logger.info("loading_docx", file=path.name)

    full_text = _read_docx(path)
    cid = _extract_candidate_id(path.name + full_text[:500])

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ".", " ", ""],
    )

    raw_doc = Document(page_content=full_text, metadata={"source": str(path)})
    chunks = splitter.split_documents([raw_doc])

    for i, chunk in enumerate(chunks):
        cid_chunk = _extract_candidate_id(chunk.page_content) or cid
        chunk.metadata.update(
            {
                "source_type": "docx",
                "source_file": path.name,
                "candidate_id": cid_chunk,
                "chunk_index": i,
            }
        )

    logger.info("docx_chunked", file=path.name, chunks=len(chunks))
    return chunks


def load_all_docx(data_dir: str = None) -> List[Document]:
    data_dir = Path(data_dir or settings.DATA_DIR)
    docs: List[Document] = []
    for docx_file in data_dir.rglob("*.docx"):
        docs.extend(load_docx(docx_file))
    logger.info("all_docx_loaded", total_chunks=len(docs))
    return docs
